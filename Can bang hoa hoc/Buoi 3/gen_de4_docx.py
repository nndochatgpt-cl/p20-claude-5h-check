# -*- coding: utf-8 -*-
"""Xuất docx format P20 (Note Level + Key) cho Đề 4 - Buổi 3, sau D5."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE = os.path.dirname(os.path.abspath(__file__))
IMG1 = os.path.join(BASE, "image", "de4_hinh1.png")
OUT = os.path.join(BASE, "Buổi 3. [SOURCE] Đề 4 - Từ hồ muối Ai Cập đến hằng số Kc.docx")

RED = RGBColor(0xEE, 0x00, 0x00)
ORANGE = RGBColor(0xFF, 0x80, 0x00)
GRAY = RGBColor(0x80, 0x80, 0x80)
GREEN = RGBColor(0x00, 0x80, 0x00)
BLUE = RGBColor(0x00, 0x70, 0xC0)
BLACK = RGBColor(0x00, 0x00, 0x00)

doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, m, Cm(2))

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), 'Times New Roman')


def set_cell_shading(cell, fill="FFFFFF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def no_border_row(table, row_idx):
    row = table.rows[row_idx]
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'nil')
            borders.append(el)
        tcPr.append(borders)


def para(text="", size=12, bold=False, italic=False, color=None, align=None, space_after=6, font=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if font:
        r.font.name = font
    return p


def multi_run_para(runs, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, kw in runs:
        r = p.add_run(text)
        r.font.size = Pt(kw.get("size", 12))
        r.bold = kw.get("bold", False)
        r.italic = kw.get("italic", False)
        if kw.get("color"):
            r.font.color.rgb = kw["color"]
    return p


def caption(label, text, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, space_before=2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r1 = p.add_run(label)
    r1.bold = True
    r1.italic = True
    r1.font.size = Pt(11)
    r2 = p.add_run(": " + text)
    r2.italic = True
    r2.font.size = Pt(11)
    return p


# ============== TIÊU ĐỀ ==============
multi_run_para([
    ("CHỦ ĐỀ 4. ", dict(bold=True, size=14, color=BLACK)),
    ("TỪ HỒ MUỐI AI CẬP ĐẾN HẰNG SỐ KC", dict(bold=True, size=14, color=RED)),
], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ============== NGỮ LIỆU ==============
ngu_lieu = [
    "[1] Trong suốt thế kỷ 18, các nhà hóa học tin rằng mỗi chất có một “ái lực” cố định với chất khác. Theo quan niệm này, một phản ứng hóa học luôn chạy hết theo một chiều duy nhất, trong đó chất có ái lực mạnh hơn sẽ chiến thắng và đẩy chất yếu hơn ra khỏi hợp chất. Không có khái niệm về sự dừng lại giữa chừng hay phản ứng quay ngược trở lại.",
    "[2] Năm 1798–1799, nhà hóa học Claude Louis Berthollet đi theo đoàn viễn chinh của Napoleon đến Ai Cập. Khi khảo sát bờ các hồ muối vùng Wadi El Natrun, ông quan sát thấy một lớp natron (Na₂CO₃) tự hình thành từ muối ăn (NaCl) trong nước hồ và đá vôi (CaCO₃) ở đáy, tức là 2NaCl + CaCO₃ ⇌ Na₂CO₃ + CaCl₂. Điều kì lạ là trong phòng thí nghiệm, phản ứng quen thuộc lại luôn chạy theo chiều ngược lại: Na₂CO₃ + CaCl₂ ⇌ CaCO₃↓ + 2NaCl.",
]
ngu_lieu2 = [
    "[3] Trong công trình Essai de statique chimique (1803), Berthollet đưa ra kết luận đột phá: khi một chất có mặt với lượng rất lớn (như NaCl cực kì đậm đặc do nước hồ bay hơi), phản ứng có thể chạy theo chiều ngược lại. Theo đó, phản ứng không nhất thiết chạy hết mà là một phản ứng thuận nghịch, có thể xảy ra theo cả chiều thuận và chiều nghịch trong cùng một điều kiện.",
]
doan4_a = "[4] Dựa trên nền tảng đó, năm 1864, Guldberg và Waage phát biểu định luật tác dụng khối lượng. Hai ông chỉ ra rằng phản ứng sẽ đạt đến trạng thái cân bằng khi tốc độ phản ứng chiều thuận bằng tốc độ phản ứng chiều nghịch. Tại đây, tỉ số giữa tích nồng độ các sản phẩm và tích nồng độ các chất đầu (mỗi nồng độ được số mũ lên bằng hệ số tỉ lượng) là một hằng số. Hằng số này gọi là hằng số cân bằng Kc. Tổng quát cho phản ứng aA + bB ⇌ cC + dD, biểu thức được viết là Kc = "
doan4_b = ". Nồng độ của chất rắn không có mặt trong biểu thức Kc."
ngu_lieu3 = [
    "[5] Vào giữa thế kỷ 20, các thí nghiệm đánh dấu đồng vị phóng xạ đã cung cấp góc nhìn vi mô. Tại trạng thái cân bằng, lượng của mỗi chất hoàn toàn không đổi, nhưng các nguyên tử vẫn liên tục di chuyển qua lại giữa chất đầu và sản phẩm. Trạng thái cân bằng hóa học do đó là một cân bằng động.",
    "[6] Về các yếu tố ảnh hưởng, Jacobus Henricus van 't Hoff (1884) chứng minh toán học rằng hằng số Kc chỉ thay đổi khi nhiệt độ của hệ thay đổi. Việc thay đổi lượng chất ban đầu, thay đổi thể tích bình chứa hay thêm chất xúc tác đều không làm hằng số Kc thay đổi.",
]

for t in ngu_lieu:
    para(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# Hình 1 ngay sau đoạn [2]
doc.add_picture(IMG1, width=Cm(13))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Hình 1", "Lớp natron kết tinh ở bờ một hồ muối vùng Wadi El Natrun (Ai Cập). Nguồn: saltworkconsultants.com", space_after=12)

for t in ngu_lieu2:
    para(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# Đoạn [4] với biểu thức Kc dùng superscript thật (không dùng dấu ^)
p4 = doc.add_paragraph()
p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p4.paragraph_format.space_after = Pt(8)
p4.add_run(doan4_a)
r = p4.add_run("[C]")
r = p4.add_run("c")
r.font.superscript = True
p4.add_run("·[D]")
r = p4.add_run("d")
r.font.superscript = True
p4.add_run(" / ([A]")
r = p4.add_run("a")
r.font.superscript = True
p4.add_run("·[B]")
r = p4.add_run("b")
r.font.superscript = True
p4.add_run(")")
p4.add_run(doan4_b)
for run in p4.runs:
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"

for t in ngu_lieu3:
    para(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

# ============== BẢNG 1 ==============
caption("Bảng 1", "Lịch sử hình thành và phát triển khái niệm cân bằng hóa học", space_after=4)

rows = [
    ("Nhà khoa học", "Năm", "Quan sát / Thí nghiệm", "Kết luận"),
    ("Berthollet", "1798–1799", "Phản ứng tạo natron tại hồ Wadi El Natrun", "Phản ứng thuận nghịch; lượng chất quyết định chiều phản ứng"),
    ("Guldberg & Waage", "1864", "Áp dụng toán học vào dữ liệu nồng độ", "Công thức tính hằng số cân bằng Kc"),
    ("(Nhiều tác giả)", "Giữa thế kỷ 20", "Thí nghiệm đánh dấu đồng vị phóng xạ", "Khẳng định bản chất của cân bằng động"),
    ("van 't Hoff", "1884", "Nghiên cứu nhiệt động lực học hóa học", "Kc chỉ phụ thuộc vào nhiệt độ"),
]
t1 = doc.add_table(rows=len(rows), cols=4)
t1.style = "Table Grid"
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = t1.cell(i, j)
        cell.text = ""
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
        r.bold = (i == 0)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Nguồn + chỉ dẫn
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("Nguồn: P20 STUDY sưu tầm và tổng hợp")
r.bold = True
r.font.size = Pt(12)
p.paragraph_format.space_after = Pt(10)

para("Đọc ngữ liệu trên và trả lời các câu hỏi sau đây:", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

# ============== CÂU HỎI ==============

def cau_label(n):
    multi_run_para([(f"Câu {n}. ", dict(bold=True, color=ORANGE, size=12)),
                     ("[P20 STUDY]", dict(bold=True, color=ORANGE, size=12))],
                    align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)


def metadata(text):
    para(text, size=10, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def noi_dung(text):
    para(text, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def option(label, text):
    para(f"{label}. {text}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def dap_an(text):
    multi_run_para([("▸ Đáp án: ", dict(bold=True, color=GREEN)), (text, dict(color=GREEN))], space_after=4)


def giai_thich_label():
    multi_run_para([("▸ Giải thích:", dict(bold=True, color=GREEN))], space_after=2)


def giai_thich_dong(text):
    para(text, color=GREEN, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def tf_table(statements):
    tbl = doc.add_table(rows=len(statements) + 1, cols=3)
    tbl.style = "Table Grid"
    widths = [Cm(12.5), Cm(1.5), Cm(1.5)]
    for i in range(len(statements) + 1):
        for j in range(3):
            tbl.cell(i, j).width = widths[j]
    # header row: no border, only 2 last cols labeled
    hcells = tbl.rows[0].cells
    hcells[0].text = ""
    set_cell_shading(hcells[0], "FFFFFF")
    for j, lbl in ((1, "Đúng"), (2, "Sai")):
        set_cell_shading(hcells[j], "FFFFFF")
        p = hcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lbl)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
    no_border_row(tbl, 0)
    for i, (letter, text) in enumerate(statements, start=1):
        c0 = tbl.cell(i, 0)
        set_cell_shading(c0, "FFFFFF")
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(f"{letter}. {text}")
        r0.font.size = Pt(11)
        r0.font.color.rgb = BLACK
        for j in (1, 2):
            set_cell_shading(tbl.cell(i, j), "FFFFFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def ghep_doi_table(cot_a, cot_b):
    n = max(len(cot_a), len(cot_b))
    tbl = doc.add_table(rows=n + 1, cols=2)
    tbl.style = "Table Grid"
    hA, hB = tbl.rows[0].cells
    for c, lbl in ((hA, "Cột A"), (hB, "Cột B")):
        set_cell_shading(c, "FFFFFF")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lbl)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
    for i in range(n):
        for col_idx, data in ((0, cot_a), (1, cot_b)):
            cell = tbl.cell(i + 1, col_idx)
            set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if i < len(data):
                r = p.add_run(data[i])
                r.font.size = Pt(11)
                r.font.color.rgb = BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def keo_tha_bank(words):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Kho từ: " + ", ".join(words))
    r.font.color.rgb = BLUE
    r.italic = False
    r.font.size = Pt(12)


# ---------- Câu 1 ----------
cau_label(1)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 1,0/5 | Mức tư duy 1 [Câu cũ số 1]")
noi_dung("Quan niệm của các nhà hóa học thế kỷ 18 về tiến trình của một phản ứng hóa học khác với góc nhìn hiện đại ở điểm nào?")
option("A", "Phản ứng hóa học chỉ dừng hẳn lại khi hằng số cân bằng bị nhiệt độ làm thay đổi.")
option("B", "Phản ứng luôn chạy đến cùng theo chiều của chất có ái lực hóa học mạnh hơn hẳn.")
option("C", "Phản ứng đạt trạng thái cân bằng khi nồng độ chất đầu bằng nồng độ sản phẩm sinh ra.")
option("D", "Phản ứng tự động ngừng hẳn khi tốc độ hai chiều thuận và nghịch trở nên bằng nhau.")
dap_an("B")
giai_thich_label()
giai_thich_dong("Căn cứ đoạn [1], thế kỷ 18 tin rằng phản ứng luôn chạy hết theo một chiều duy nhất, trong đó chất có ái lực mạnh hơn sẽ chiến thắng.")
giai_thich_dong("A sai lịch sử: đây là kết luận của van 't Hoff (đoạn [6]), không phải quan niệm thế kỷ 18.")
giai_thich_dong("C mắc lỗi L2: đánh đồng trạng thái cân bằng với việc nồng độ hai bên bằng nhau.")
giai_thich_dong("D mắc lỗi L1: cho rằng phản ứng ngừng hẳn, trái với bản chất cân bằng động ở đoạn [5].")

# ---------- Câu 2 ----------
cau_label(2)
metadata("Đúng–Sai | Độ khó 1,5/5 | Mức tư duy 1 [Câu cũ số 2]")
noi_dung("Các nhận định sau đúng hay sai?")
tf_table([
    ("A", "Berthollet kết luận rằng phản ứng hóa học không bắt buộc phải luôn diễn ra đến khi hết chất đầu."),
    ("B", "Thí nghiệm đánh dấu đồng vị cho thấy khi hệ đạt cân bằng, phản ứng ở cả hai chiều đều dừng hẳn lại."),
    ("C", "Hằng số cân bằng Kc có giá trị càng lớn thì phản ứng thuận nghịch càng nhanh đạt tới trạng thái cân bằng."),
    ("D", "Theo Guldberg và Waage, cân bằng được thiết lập khi tốc độ chiều thuận bằng tốc độ chiều nghịch."),
])
dap_an("Đ, S, S, Đ")
giai_thich_label()
giai_thich_dong("A đúng: căn cứ đoạn [3], phản ứng không nhất thiết chạy hết.")
giai_thich_dong("B sai (L1): đoạn [5] khẳng định nguyên tử vẫn liên tục chuyển, đó là cân bằng động, không phải dừng hẳn.")
giai_thich_dong("C sai (L5): Kc là tỉ số nồng độ lúc cân bằng, không quyết định tốc độ đạt tới cân bằng.")
giai_thich_dong("D đúng: định nghĩa trực tiếp từ đoạn [4].")

# ---------- Câu 3 ----------
cau_label(3)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 2,5/5 | Mức tư duy 2 [Câu cũ số 5]")
noi_dung("Điều kiện nào sau đây sẽ làm thay đổi giá trị của hằng số cân bằng Kc đối với một hệ phản ứng đang ở trạng thái cân bằng trong bình kín?")
option("A", "Bơm thêm một lượng lớn chất phản ứng vào bình chứa hệ đang cân bằng.")
option("B", "Giảm thể tích của bình chứa hệ đang cân bằng để làm tăng áp suất.")
option("C", "Điều chỉnh lượng chất nạp vào sao cho nồng độ chất đầu bằng sản phẩm.")
option("D", "Đưa bình chứa hệ phản ứng đang cân bằng vào trong một bể nước đá.")
dap_an("D")
giai_thich_label()
giai_thich_dong("Căn cứ đoạn [6], Kc chỉ thay đổi khi nhiệt độ thay đổi. Đưa vào bể nước đá làm thay đổi nhiệt độ của hệ.")
giai_thich_dong("A và B mắc lỗi L3: thêm chất hoặc đổi thể tích không làm đổi Kc.")
giai_thich_dong("C mắc lỗi L2: không có yêu cầu nồng độ hai bên phải bằng nhau.")

# ---------- Câu 4 ----------
cau_label(4)
metadata("Đúng–Sai | Độ khó 2,5/5 | Mức tư duy 2 [Câu cũ số 6]")
noi_dung("Các nhận định sau đúng hay sai?")
tf_table([
    ("A", "Tại trạng thái cân bằng, nồng độ của chất đầu bắt buộc phải bằng nồng độ của sản phẩm sinh ra."),
    ("B", "Nguyên tử của các chất vẫn liên tục tham gia phản ứng hai chiều ngay cả khi hệ đã đạt cân bằng."),
    ("C", "Việc bổ sung chất xúc tác vào hệ phản ứng không làm thay đổi giá trị của hằng số cân bằng Kc."),
    ("D", "Làm cho bình phản ứng giãn nở thể tích sẽ làm thay đổi giá trị hằng số cân bằng Kc của hệ."),
])
dap_an("S, Đ, Đ, S")
giai_thich_label()
giai_thich_dong("A sai (L2): tại cân bằng nồng độ ổn định, không bắt buộc bằng nhau.")
giai_thich_dong("B đúng: cân bằng động, căn cứ đoạn [5].")
giai_thich_dong("C đúng: căn cứ kết luận van 't Hoff, đoạn [6].")
giai_thich_dong("D sai (L3): thể tích không làm đổi Kc, chỉ nhiệt độ mới làm đổi.")

# ---------- Câu 5 ----------
cau_label(5)
metadata("Điền khuyết | Độ khó 2,0/5 | Mức tư duy 2 [Câu cũ số 7]")
noi_dung("Điền từ/cụm từ (không quá 2 tiếng) hoặc số (làm tròn đến hàng đơn vị) vào chỗ trống:")
noi_dung("Một hệ hóa học phản ứng theo phương trình X(g) + 2Y(g) ⇌ Z(g). Tại trạng thái cân bằng ở một nhiệt độ nhất định, đo được nồng độ các chất là [X] = 0,10 M, [Y] = 0,20 M, và [Z] = 0,08 M. Giá trị hằng số Kc của phản ứng này là ..........")
dap_an("20")
giai_thich_label()
giai_thich_dong("Áp dụng công thức đoạn [4]: Kc = [Z]/([X]·[Y]²) = 0,08/(0,10 × 0,20²) = 0,08/0,004 = 20.")

# ---------- Câu 6 ----------
cau_label(6)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 3,5/5 | Mức tư duy 3 | Giả định tình huống [Câu cũ số 9]")
noi_dung("Áp dụng lập luận của Berthollet, muốn phản ứng BaSO₄(s) + 2NaCl(aq) ⇌ BaCl₂(aq) + Na₂SO₄(aq) (vốn rất khó xảy ra) chuyển dịch mạnh theo chiều tạo ra các muối tan ở một hồ nước tự nhiên, điều kiện nào sau đây là phù hợp nhất?")
option("A", "Hồ nước phải cạn dần do nắng nóng để tạo ra nồng độ muối ăn cực kì lớn.")
option("B", "Hồ nước cần được pha loãng để nồng độ muối ăn bằng nồng độ bari clorua.")
option("C", "Phản ứng cần có hằng số Kc vô cùng lớn để tốc độ tạo muối tan diễn ra tức thì.")
option("D", "Hồ nước phải đứng yên để phản ứng thuận ngừng hẳn lại tại điểm cân bằng.")
dap_an("A")
giai_thich_label()
giai_thich_dong("Berthollet kết luận lượng chất cực lớn (như NaCl đậm đặc do bay hơi) ép phản ứng chạy theo chiều tạo thêm sản phẩm, giống cơ chế tạo natron ở đoạn [2]–[3].")
giai_thich_dong("B mắc lỗi L2. C mắc lỗi L5. D mắc lỗi L1.")

# ---------- Câu 7 ----------
cau_label(7)
metadata("Ghép đôi | Độ khó 3,5/5 | Mức tư duy 3 | Giả định tình huống [Câu cũ số 12]")
noi_dung("Trong quá trình nghiên cứu phản ứng 2E(g) ⇌ F(g), nhà hóa học thực hiện các thao tác ở Cột A. Ghép mỗi ý ở Cột A với một ý ở Cột B.")
ghep_doi_table(
    ["1. Đưa toàn bộ bình phản ứng vào hầm có nhiệt độ cực thấp.",
     "2. Bơm thêm khí E vào trong bình chứa (giữ nguyên nhiệt độ).",
     "3. Đưa chất đồng vị phóng xạ E* vào hệ đang cân bằng.",
     "4. Đo đạc nồng độ các chất ngay ở giây đầu tiên phản ứng."],
    ["a. Làm thay đổi trực tiếp giá trị của hằng số cân bằng Kc trong hệ.",
     "b. Làm thay đổi trạng thái cân bằng nhưng giá trị Kc không hề đổi theo.",
     "c. Chứng minh thực nghiệm nguyên tử liên tục trao đổi qua lại hai chiều.",
     "d. Cung cấp ngay một tỉ số nồng độ chưa bằng với giá trị Kc của hệ."],
)
dap_an("1-a, 2-b, 3-c, 4-d")
giai_thich_label()
giai_thich_dong("1-a: nhiệt độ làm đổi Kc (đoạn [6]).")
giai_thich_dong("2-b: thêm chất ép cân bằng chuyển dịch nhưng Kc cố định, bác bỏ L3.")
giai_thich_dong("3-c: thí nghiệm đồng vị chứng minh cân bằng động (đoạn [5]).")
giai_thich_dong("4-d: tỉ số nồng độ chỉ bằng Kc khi hệ đã đạt cân bằng.")

# ---------- Câu 8 ----------
cau_label(8)
metadata("Ghép đôi | Độ khó 4,0/5 | Mức tư duy 4 | Chẩn đoán & Khắc phục [Câu cũ số 13]")
noi_dung("Ghép các phát biểu sai lầm của một sinh viên hóa học ở Cột A với lỗ hổng tư duy dẫn đến sai lầm đó ở Cột B.")
ghep_doi_table(
    ["1. “Màu của bình chứa không đổi nữa nên các chất khí bên trong đã ngừng phản ứng với nhau.”",
     "2. “Để phản ứng tạo ra trạng thái cân bằng, ta phải nạp đúng lượng sao cho chất đầu bằng sản phẩm.”",
     "3. “Phản ứng này có Kc = 10⁸ nên chắc chắn chỉ cần tích tắc là quá trình hóa học hoàn tất.”",
     "4. “Đập vụn chất rắn nạp thêm vào bình kín sẽ làm giá trị hằng số cân bằng Kc tăng lên.”"],
    ["a. Đánh đồng tính vô hướng của nhiệt động học với tốc độ của động học.",
     "b. Không phân biệt được sự bằng nhau của tốc độ hai chiều với nồng độ.",
     "c. Lãng quên chất rắn bị loại khỏi biểu thức Kc, chỉ phụ thuộc nhiệt độ.",
     "d. Phủ nhận bản chất động học liên tục không ngừng nghỉ của cân bằng hệ."],
)
dap_an("1-d, 2-b, 3-a, 4-c")
giai_thich_label()
giai_thich_dong("1-d (lỗi L1): tưởng cân bằng là trạng thái tĩnh.")
giai_thich_dong("2-b (lỗi L2): nhầm nồng độ bằng nhau với tốc độ hai chiều bằng nhau.")
giai_thich_dong("3-a (lỗi L5): Kc lớn không nói lên tốc độ phản ứng.")
giai_thich_dong("4-c (lỗi L3, L4): chất rắn không có mặt trong Kc; Kc chỉ phụ thuộc nhiệt độ.")

# ---------- Câu 9 ----------
cau_label(9)
metadata("Kéo thả | Độ khó 3,5/5 | Mức tư duy 3 | Mô hình hóa & Tỷ lệ [Câu cũ số 14]")
noi_dung("Kéo các cụm từ trong kho vào chỗ trống cho phù hợp:")
keo_tha_bank(["NaCl", "MSO₄", "MCl₂", "2", "1", "thuận", "nghịch"])
noi_dung("Xét một phản ứng giả định sinh ra khoáng chất mới: MCl₂(aq) + Na₂SO₄(aq) ⇌ MSO₄(s) + 2NaCl(aq). Dựa theo định luật tác dụng khối lượng, biểu thức Kc của phản ứng này sẽ có dạng một phân thức, trong đó nồng độ (1)______ nằm ở tử số với số mũ bằng (2)______. Do quy ước toán học của hệ dị thể, nồng độ của chất (3)______ không xuất hiện trong biểu thức này. Vận dụng trực tiếp lập luận của Berthollet, nếu nước hồ tự nhiên bay hơi khiến nồng độ muối ăn trở nên cực kì đậm đặc, hệ thống sẽ bị ép buộc phản ứng ưu tiên theo chiều (4)______.")
dap_an("(1) NaCl  (2) 2  (3) MSO₄  (4) nghịch")
giai_thich_label()
giai_thich_dong("Kc = [NaCl]²/([MCl₂][Na₂SO₄]). MSO₄ là chất rắn, không vào biểu thức (bác bỏ L4). NaCl là sản phẩm; dư thừa NaCl ép hệ chạy ngược lại để tiêu bớt NaCl, tức chiều nghịch.")

# ============== BẢNG TỔNG HỢP ==============
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BẢNG TỔNG HỢP")
r.bold = True
r.font.size = Pt(13)
p.paragraph_format.space_after = Pt(8)

summary = [
    ("Câu", "Dạng", "Độ khó", "Mức", "Kỹ năng nâng cao", "Nguồn", "Đáp án"),
    ("1", "MCQ", "1,0", "1", "—", "Cũ #1", "B"),
    ("2", "Đúng–Sai", "1,5", "1", "—", "Cũ #2", "Đ,S,S,Đ"),
    ("3", "MCQ", "2,5", "2", "—", "Cũ #5", "D"),
    ("4", "Đúng–Sai", "2,5", "2", "—", "Cũ #6", "S,Đ,Đ,S"),
    ("5", "Điền khuyết", "2,0", "2", "—", "Cũ #7", "20"),
    ("6", "MCQ", "3,5", "3", "Giả định tình huống", "Cũ #9", "A"),
    ("7", "Ghép đôi", "3,5", "3", "Giả định tình huống", "Cũ #12", "1a,2b,3c,4d"),
    ("8", "Ghép đôi", "4,0", "4", "Chẩn đoán & Khắc phục", "Cũ #13", "1d,2b,3a,4c"),
    ("9", "Kéo thả", "3,5", "3", "Mô hình hóa & Tỷ lệ", "Cũ #14", "NaCl,2,MSO₄,nghịch"),
]
t2 = doc.add_table(rows=len(summary), cols=7)
t2.style = "Table Grid"
widths = [Cm(1.0), Cm(2.2), Cm(1.6), Cm(1.2), Cm(3.6), Cm(1.8), Cm(3.6)]
for i, row in enumerate(summary):
    for j, val in enumerate(row):
        cell = t2.cell(i, j)
        cell.width = widths[j]
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.font.size = Pt(10)
        r.font.color.rgb = BLACK
        r.bold = (i == 0)

doc.save(OUT)
print("Saved:", OUT)
