# -*- coding: utf-8 -*-
"""Xuất docx format P20 (Note Level + Key) cho Đề 5 - Buổi 3, sau D4 (10 câu đã lọc, chưa D5)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "Buổi 3. [D5-sơ bộ] Đề 5 - Nhật ký lò tổng hợp methanol.docx")

RED = RGBColor(0xEE, 0x00, 0x00)
ORANGE = RGBColor(0xFF, 0x80, 0x00)
GRAY = RGBColor(0x80, 0x80, 0x80)
GREEN = RGBColor(0x00, 0x80, 0x00)
BLUE = RGBColor(0x00, 0x70, 0xC0)
BLACK = RGBColor(0x00, 0x00, 0x00)

doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, m, Cm(2))

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), 'Times New Roman')


def set_cell_shading(cell, fill="FFFFFF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def no_border_row(table, row_idx):
    row = table.rows[row_idx]
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'nil')
            borders.append(el)
        tcPr.append(borders)


def para(text="", size=12, bold=False, italic=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def multi_run_para(runs, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, kw in runs:
        r = p.add_run(text)
        r.font.size = Pt(kw.get("size", 12))
        r.bold = kw.get("bold", False)
        r.italic = kw.get("italic", False)
        if kw.get("color"):
            r.font.color.rgb = kw["color"]
    return p


def caption(label, text, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r1 = p.add_run(label)
    r1.bold = True
    r1.italic = True
    r1.font.size = Pt(11)
    r2 = p.add_run(": " + text)
    r2.italic = True
    r2.font.size = Pt(11)
    return p


# ============== TIÊU ĐỀ ==============
multi_run_para([
    ("CHỦ ĐỀ 5. ", dict(bold=True, size=14, color=BLACK)),
    ("NHẬT KÝ LÒ TỔNG HỢP METHANOL", dict(bold=True, size=14, color=RED)),
], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = note.add_run("[D5 sơ bộ trên 10 câu đã lọc D4 — GĐ1 nội dung và GĐ2 hình thức đã kiểm, chưa có Hình 1. Sẽ chạy D5 lại sau khi bạn chốt final list.]")
r.italic = True
r.font.size = Pt(10)
r.font.color.rgb = GRAY
note.paragraph_format.space_after = Pt(14)

# ============== NGỮ LIỆU ==============
ngu_lieu = [
    "[1] Methanol (CH₃OH) là nguyên liệu và nhiên liệu công nghiệp quan trọng. Phản ứng tổng hợp methanol từ khí than được thực hiện theo phương trình hóa học: CO(g) + 2H₂(g) ⇌ CH₃OH(g). Phản ứng này diễn ra theo chiều thuận tỏa nhiệt và sử dụng hệ xúc tác Cu/ZnO.",
    "[2] Tại trạng thái cân bằng, hằng số nồng độ Kc của hệ được tính bằng biểu thức Kc = [CH₃OH]/([CO][H₂]²). Giá trị Kc của một phản ứng xác định chỉ phụ thuộc vào nhiệt độ. Khi hệ đang ở cân bằng mà bị tác động bởi các yếu tố bên ngoài, nồng độ các chất có thể thay đổi, làm xuất hiện thương số phản ứng Q. Thương số Q được tính bằng công thức tương tự Kc nhưng áp dụng cho nồng độ tại thời điểm đang xét. Việc đối chiếu Q và Kc giúp xác định chiều dịch chuyển của phản ứng: hệ chạy theo chiều thuận nếu Q nhỏ hơn Kc; hệ chạy theo chiều nghịch nếu Q lớn hơn Kc; hệ đạt cân bằng khi Q bằng Kc. Mọi thao tác làm thay đổi nồng độ hoặc thể tích chỉ tác động lên Q, hệ phản ứng sau đó sẽ tự dịch chuyển để đưa Q quay về đúng giá trị Kc ban đầu. Chỉ tác động làm thay đổi nhiệt độ mới đưa hệ chuyển dời tới một giá trị Kc mới.",
    "[3] Một lò phản ứng kín dung tích V lít đang duy trì ổn định ở trạng thái cân bằng tại 250 °C (gọi là trạng thái gốc). Tại trạng thái gốc này, nồng độ các chất đo được là: [CO] = 0,50 M; [H₂] = 1,0 M; [CH₃OH] = 0,25 M. Để khảo sát động học, kỹ thuật viên tiến hành 5 thao tác thử nghiệm độc lập. Mỗi thao tác đều xuất phát từ trạng thái gốc và không nối tiếp nhau. Thiết bị cảm biến ghi nhận lại nồng độ các chất ngay tức thì sau thao tác, khi phản ứng chưa kịp dịch chuyển bù trừ. Dữ liệu được ghi nhận chi tiết trong Bảng 1. Cùng với đó, Bảng 2 cung cấp giá trị Kc chuẩn đo được tại các mốc nhiệt độ khác nhau.",
]
for t in ngu_lieu:
    para(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

caption("Bảng 1", "Nhật ký 5 thao tác độc lập")
rows1 = [
    ("Thao tác", "Nhiệt độ (°C)", "[CO] ngay sau thao tác (M)", "[H₂] ngay sau thao tác (M)", "[CH₃OH] ngay sau thao tác (M)"),
    ("1. Bơm thêm CO", "250", "1,00", "1,0", "0,25"),
    ("2. Rút bớt CH₃OH", "250", "0,50", "1,0", "0,125"),
    ("3. Nén giảm thể tích còn một nửa", "250", "1,00", "2,0", "0,50"),
    ("4. Nâng nhiệt độ", "300", "0,50", "1,0", "0,25"),
    ("5. Thay lớp xúc tác mới", "250", "0,50", "1,0", "0,25"),
]
t1 = doc.add_table(rows=len(rows1), cols=5)
t1.style = "Table Grid"
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
widths1 = [Cm(3.6), Cm(2.0), Cm(2.6), Cm(2.6), Cm(2.6)]
for i, row in enumerate(rows1):
    for j, val in enumerate(row):
        cell = t1.cell(i, j)
        cell.width = widths1[j]
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        r.font.size = Pt(10.5)
        r.font.color.rgb = BLACK
        r.bold = (i == 0)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

caption("Bảng 2", "Giá trị Kc theo nhiệt độ")
rows2 = [
    ("Nhiệt độ (°C)", "Kc"),
    ("250", "0,50"),
    ("300", "0,20"),
]
t2 = doc.add_table(rows=len(rows2), cols=2)
t2.style = "Table Grid"
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
widths2 = [Cm(4), Cm(3)]
for i, row in enumerate(rows2):
    for j, val in enumerate(row):
        cell = t2.cell(i, j)
        cell.width = widths2[j]
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.font.size = Pt(10.5)
        r.font.color.rgb = BLACK
        r.bold = (i == 0)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("Nguồn: P20 STUDY sưu tầm và tổng hợp")
r.bold = True
r.font.size = Pt(12)
p.paragraph_format.space_after = Pt(10)

para("Đọc ngữ liệu trên và trả lời các câu hỏi sau đây:", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

# ============== CÂU HỎI ==============

def cau_label(n):
    multi_run_para([(f"Câu {n}. ", dict(bold=True, color=ORANGE, size=12)),
                     ("[P20 STUDY]", dict(bold=True, color=ORANGE, size=12))],
                    align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)


def metadata(text):
    para(text, size=10, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def noi_dung(text):
    para(text, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def option(label, text):
    para(f"{label}. {text}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def dap_an(text):
    multi_run_para([("▸ Đáp án: ", dict(bold=True, color=GREEN)), (text, dict(color=GREEN))], space_after=4)


def giai_thich_label():
    multi_run_para([("▸ Giải thích:", dict(bold=True, color=GREEN))], space_after=2)


def giai_thich_dong(text):
    para(text, color=GREEN, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)


def tf_table(statements):
    tbl = doc.add_table(rows=len(statements) + 1, cols=3)
    tbl.style = "Table Grid"
    widths = [Cm(12.5), Cm(1.5), Cm(1.5)]
    for i in range(len(statements) + 1):
        for j in range(3):
            tbl.cell(i, j).width = widths[j]
    hcells = tbl.rows[0].cells
    hcells[0].text = ""
    set_cell_shading(hcells[0], "FFFFFF")
    for j, lbl in ((1, "Đúng"), (2, "Sai")):
        set_cell_shading(hcells[j], "FFFFFF")
        p = hcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lbl)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
    no_border_row(tbl, 0)
    for i, (letter, text) in enumerate(statements, start=1):
        c0 = tbl.cell(i, 0)
        set_cell_shading(c0, "FFFFFF")
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(f"{letter}. {text}")
        r0.font.size = Pt(11)
        r0.font.color.rgb = BLACK
        for j in (1, 2):
            set_cell_shading(tbl.cell(i, j), "FFFFFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def ghep_doi_table(cot_a, cot_b):
    n = max(len(cot_a), len(cot_b))
    tbl = doc.add_table(rows=n + 1, cols=2)
    tbl.style = "Table Grid"
    hA, hB = tbl.rows[0].cells
    for c, lbl in ((hA, "Cột A"), (hB, "Cột B")):
        set_cell_shading(c, "FFFFFF")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lbl)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
    for i in range(n):
        for col_idx, data in ((0, cot_a), (1, cot_b)):
            cell = tbl.cell(i + 1, col_idx)
            set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if i < len(data):
                r = p.add_run(data[i])
                r.font.size = Pt(11)
                r.font.color.rgb = BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def keo_tha_bank(words):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Kho từ: " + ", ".join(words))
    r.font.color.rgb = BLUE
    r.font.size = Pt(12)


# ---------- Câu 1 ----------
cau_label(1)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 1,0/5 | Mức tư duy 1 [Câu cũ số 1]")
noi_dung("Dựa vào ngữ liệu, thao tác nào dưới đây làm thay đổi hằng số cân bằng Kc?")
option("A", "Thao tác 1 (Bơm thêm một lượng khí CO vào lò).")
option("B", "Thao tác 3 (Nén giảm thể tích bình phản ứng).")
option("C", "Thao tác 4 (Nâng nhiệt độ của lò phản ứng).")
option("D", "Thao tác 5 (Thay lớp xúc tác mới trong lò).")
dap_an("C")
giai_thich_label()
giai_thich_dong("Đoạn [2]: chỉ nhiệt độ mới làm đổi Kc. A, B, D mắc lỗi L1 (tưởng thêm chất, đổi áp suất, đổi xúc tác làm đổi Kc).")

# ---------- Câu 2 ----------
cau_label(2)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 2,0/5 | Mức tư duy 2 [Câu cũ số 2]")
noi_dung("Đối với thao tác 1, nhận định nào mô tả đúng kết quả so sánh và chiều dịch chuyển?")
option("A", "Q = 0,25; do đó phản ứng sẽ dịch chuyển theo chiều thuận.")
option("B", "Q = 0,25; do đó phản ứng sẽ dịch chuyển theo chiều nghịch.")
option("C", "Q = 1,00; do đó phản ứng sẽ dịch chuyển theo chiều nghịch.")
option("D", "Q = 0,50; do đó hệ tiếp tục duy trì trạng thái cân bằng cũ.")
dap_an("A")
giai_thich_label()
giai_thich_dong("Q = 0,25/(1,00×1,0²) = 0,25 < Kc = 0,50, nên hệ dịch chuyển theo chiều thuận. B mắc lỗi L5 (nhầm chiều). C sai tính toán. D mắc lỗi L1.")

# ---------- Câu 3 ----------
cau_label(3)
metadata("Trắc nghiệm 4 lựa chọn | Độ khó 3,5/5 | Mức tư duy 3 | Bài toán ngược & Điều kiện biên [Câu cũ số 4]")
noi_dung("Giả sử kỹ thuật viên tiến hành một thao tác giãn thể tích lò phản ứng lên gấp đôi (giữ nguyên 250 °C) từ trạng thái gốc. Hệ quả đối với hệ lúc này là gì?")
option("A", "Q = 2,0; lớn hơn Kc nên phản ứng dịch chuyển theo chiều nghịch.")
option("B", "Q = 2,0; lớn hơn Kc nên phản ứng dịch chuyển theo chiều thuận.")
option("C", "Q = 0,5; bằng Kc nên hệ tiếp tục duy trì trạng thái cân bằng cũ.")
option("D", "Q = 0,125; nhỏ hơn Kc nên phản ứng dịch chuyển theo chiều thuận.")
dap_an("A")
giai_thich_label()
giai_thich_dong("Giãn gấp đôi làm mọi nồng độ giảm một nửa: [CO]=0,25; [H₂]=0,50; [CH₃OH]=0,125. Q = 0,125/(0,25×0,50²) = 2,0 > Kc = 0,50, nên hệ dịch chuyển theo chiều nghịch.")
giai_thich_dong("B mắc lỗi L5. C mắc lỗi L1. D sai tính toán.")

# ---------- Câu 4 ----------
cau_label(4)
metadata("Đúng–Sai | Độ khó 1,5/5 | Mức tư duy 1 [Câu cũ số 5]")
noi_dung("Các nhận định sau đúng hay sai khi phân tích thao tác 5?")
tf_table([
    ("A", "Thao tác 5 làm biến đổi giá trị hằng số cân bằng Kc của hệ."),
    ("B", "Ngay sau thao tác 5, thương số phản ứng Q vẫn giữ giá trị 0,50."),
    ("C", "Lớp xúc tác mới làm tăng nồng độ CH₃OH ở trạng thái cân bằng."),
    ("D", "Hệ phản ứng không dịch chuyển, chỉ rút ngắn thời gian cân bằng."),
])
dap_an("S, Đ, S, Đ")
giai_thich_label()
giai_thich_dong("A sai (L1): chỉ nhiệt độ đổi Kc. B đúng: nồng độ không đổi nên Q = 0,50.")
giai_thich_dong("C sai (L4): xúc tác không đổi vị trí cân bằng. D đúng: Q = Kc nên không dịch chuyển, xúc tác chỉ tăng tốc độ.")

# ---------- Câu 5 ----------
cau_label(5)
metadata("Đúng–Sai | Độ khó 2,5/5 | Mức tư duy 2 [Câu cũ số 6]")
noi_dung("Các nhận định sau đúng hay sai khi phân tích thao tác 3?")
tf_table([
    ("A", "Ngay sau khi nén thể tích, thương số phản ứng Q có giá trị bằng 0,125."),
    ("B", "Do giá trị Q nhỏ hơn Kc, hệ phản ứng sẽ chuyển dịch theo chiều nghịch."),
    ("C", "Lượng phân tử khí trong bình chứa sẽ giảm dần trong quá trình dịch chuyển."),
    ("D", "Sau khi đạt cân bằng mới, nồng độ các chất lập tức quay về giá trị ban đầu."),
])
dap_an("Đ, S, Đ, S")
giai_thich_label()
giai_thich_dong("A đúng: Q = 0,50/(1,00×2,0²) = 0,125. B sai (L5): Q<Kc phải là chiều thuận.")
giai_thich_dong("C đúng: chiều thuận làm giảm số mol khí, từ 3 mol chất đầu còn 1 mol sản phẩm. D sai (L3): chỉ bù một phần, không quay hẳn về giá trị cũ.")

# ---------- Câu 6 ----------
cau_label(6)
metadata("Điền khuyết | Độ khó 2,0/5 | Mức tư duy 2 [Câu cũ số 8]")
noi_dung("Điền từ/cụm từ (không quá 2 tiếng) hoặc số (làm tròn đến 2 chữ số thập phân) vào chỗ trống:")
noi_dung("Ngay sau thao tác 2, thương số phản ứng Q có giá trị là (1)______, dẫn đến kết quả hệ bắt buộc phải dịch chuyển theo chiều (2)______.")
dap_an("(1) 0,25  (2) thuận")
giai_thich_label()
giai_thich_dong("Q = 0,125/(0,50×1,0²) = 0,25 < Kc = 0,50, nên chiều thuận.")

# ---------- Câu 7 ----------
cau_label(7)
metadata("Điền khuyết | Độ khó 4,5/5 | Mức tư duy 4 | Bài toán ngược & Điều kiện biên [Câu cũ số 10]")
noi_dung("Điền từ/cụm từ (không quá 2 tiếng) hoặc số (làm tròn đến 2 chữ số thập phân) vào chỗ trống:")
noi_dung("Thực hiện một thao tác phức hợp từ trạng thái gốc: rút lượng H₂ đi một nửa đồng thời nâng nhiệt độ lò lên 300 °C. Ngay lúc này, thương số Q tính được có giá trị là (1)______ và phản ứng bắt đầu dịch chuyển để hướng tới trạng thái có hằng số cân bằng bằng (2)______.")
dap_an("(1) 2,00  (2) 0,20")
giai_thich_label()
giai_thich_dong("[H₂]=0,50 M (rút nửa); [CO], [CH₃OH] chưa đổi. Q = 0,25/(0,50×0,50²) = 2,00. Nhiệt độ lên 300 °C nên đích là Kc(300°C) = 0,20.")

# ---------- Câu 8 ----------
cau_label(8)
metadata("Ghép đôi | Độ khó 3,5/5 | Mức tư duy 3 | Mô hình hóa & Tỷ lệ [Câu cũ số 12]")
noi_dung("Ghép mỗi ý ở Cột A với một ý tương ứng ở Cột B về mối quan hệ tỉ lệ giữa Q và Kc khi thực hiện các tác động nồng độ độc lập tại 250 °C.")
ghep_doi_table(
    ["1. Nén giảm thể tích bình chứa còn một nửa.",
     "2. Bơm thêm khí CO để nồng độ CO tăng gấp đôi.",
     "3. Giãn thể tích bình chứa lên gấp đôi.",
     "4. Bơm thêm CH₃OH để nồng độ CH₃OH tăng gấp đôi."],
    ["a. Q ngay sau thao tác bằng 0,25 lần Kc.",
     "b. Q ngay sau thao tác bằng 0,50 lần Kc.",
     "c. Q ngay sau thao tác bằng 4,00 lần Kc.",
     "d. Q ngay sau thao tác bằng 2,00 lần Kc."],
)
dap_an("1-a, 2-b, 3-c, 4-d")
giai_thich_label()
giai_thich_dong("(1) Nén 1/2 làm mọi nồng độ gấp đôi, do đó Q mới bằng Q cũ chia 4, tức 0,25Kc. (2) CO gấp đôi ở mẫu số nên Q bằng 0,50Kc.")
giai_thich_dong("(3) Giãn gấp đôi làm mọi nồng độ giảm nửa, do đó Q mới bằng Q cũ nhân 4, tức 4,00Kc. (4) CH₃OH gấp đôi ở tử số nên Q bằng 2,00Kc.")

# ---------- Câu 9 ----------
cau_label(9)
metadata("Ghép đôi | Độ khó 4,0/5 | Mức tư duy 4 | Phân tích dữ liệu chéo [Câu cũ số 13]")
noi_dung("Ghép mỗi ý ở Cột A (các thao tác xuất phát từ trạng thái gốc) với một kết quả logic ở Cột B (đợi hệ đạt cân bằng mới).")
ghep_doi_table(
    ["1. Thao tác 1: bơm thêm CO vào hệ ở 250 °C.",
     "2. Thao tác 4: nâng nhiệt độ lò lên 300 °C.",
     "3. Thao tác 5: thay xúc tác mới có hoạt tính cao.",
     "4. Thao tác 2: rút bớt lượng CH₃OH ra khỏi hệ ở 250 °C."],
    ["a. Nồng độ CH₃OH lúc hệ đạt cân bằng mới cao hơn mức 0,25 M ban đầu.",
     "b. Nồng độ CH₃OH lúc hệ đạt cân bằng mới thấp hơn mức 0,25 M ban đầu.",
     "c. Nồng độ CH₃OH lúc hệ đạt cân bằng mới vẫn giữ đúng mức 0,25 M.",
     "d. Nồng độ CH₃OH lúc hệ đạt cân bằng mới nằm khoảng 0,125 đến 0,25 M."],
)
dap_an("1-a, 2-b, 3-c, 4-d")
giai_thich_label()
giai_thich_dong("(1) Thêm CO làm hệ chạy thuận, CH₃OH tăng so với 0,25 gốc. (2) Tăng nhiệt làm hệ chạy nghịch, CH₃OH giảm dưới 0,25.")
giai_thich_dong("(3) Xúc tác giữ nguyên mức 0,25. (4) Rút CH₃OH xuống 0,125 khiến hệ chạy thuận nhưng chỉ bù một phần (L3), không bao giờ đạt lại 0,25.")

# ---------- Câu 10 ----------
cau_label(10)
metadata("Kéo thả | Độ khó 4,0/5 | Mức tư duy 4 | Chẩn đoán & Khắc phục [Câu cũ số 15]")
noi_dung("Kéo các cụm từ trong kho vào chỗ trống cho phù hợp:")
keo_tha_bank(["giảm xuống", "tăng lên", "cao hơn", "thấp hơn", "bằng với"])
noi_dung("Khi thực hiện thao tác 1 (thêm CO), phản ứng dịch chuyển theo chiều thuận làm lượng CO (1)______ so với thời điểm ngay sau thao tác. Theo nguyên lý Le Chatelier, sự dịch chuyển này chỉ bù đắp một phần tác động, dẫn đến nồng độ CO lúc đạt cân bằng mới rốt cuộc vẫn (2)______ so với nồng độ ở trạng thái gốc ban đầu (0,50 M).")
dap_an("(1) giảm xuống  (2) cao hơn")
giai_thich_label()
giai_thich_dong("CO vọt lên 1,00 M rồi giảm dần theo chiều thuận. Vì chỉ bù một phần (L3), CO cuối cùng dừng ở mức giữa 0,50 và 1,00 M, tức cao hơn trạng thái gốc.")

# ============== BẢNG TỔNG HỢP ==============
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BẢNG TỔNG HỢP")
r.bold = True
r.font.size = Pt(13)
p.paragraph_format.space_after = Pt(8)

summary = [
    ("Câu", "Dạng", "Độ khó", "Mức", "Kỹ năng nâng cao", "Nguồn", "Đáp án"),
    ("1", "MCQ", "1,0", "1", "—", "Cũ #1", "C"),
    ("2", "MCQ", "2,0", "2", "—", "Cũ #2", "A"),
    ("3", "MCQ", "3,5", "3", "BT ngược & ĐK biên", "Cũ #4", "A"),
    ("4", "Đúng–Sai", "1,5", "1", "—", "Cũ #5", "S,Đ,S,Đ"),
    ("5", "Đúng–Sai", "2,5", "2", "—", "Cũ #6", "Đ,S,Đ,S"),
    ("6", "Điền khuyết", "2,0", "2", "—", "Cũ #8", "0,25 / thuận"),
    ("7", "Điền khuyết", "4,5", "4", "BT ngược & ĐK biên", "Cũ #10", "2,00 / 0,20"),
    ("8", "Ghép đôi", "3,5", "3", "Mô hình hóa & Tỷ lệ", "Cũ #12", "1a,2b,3c,4d"),
    ("9", "Ghép đôi", "4,0", "4", "Phân tích dữ liệu chéo", "Cũ #13", "1a,2b,3c,4d"),
    ("10", "Kéo thả", "4,0", "4", "Chẩn đoán & Khắc phục", "Cũ #15", "giảm xuống / cao hơn"),
]
t3 = doc.add_table(rows=len(summary), cols=7)
t3.style = "Table Grid"
widths = [Cm(1.0), Cm(2.0), Cm(1.6), Cm(1.2), Cm(3.6), Cm(1.8), Cm(3.8)]
for i, row in enumerate(summary):
    for j, val in enumerate(row):
        cell = t3.cell(i, j)
        cell.width = widths[j]
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.font.size = Pt(10)
        r.font.color.rgb = BLACK
        r.bold = (i == 0)

doc.save(OUT)
print("Saved:", OUT)
